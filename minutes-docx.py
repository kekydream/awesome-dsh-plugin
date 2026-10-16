"""Render one meeting's minutes as a Word document.

本插件由博尔塔拉融媒体中心媒体传播与认知重点实验室制作。
制作人：杨道泉　邮箱：kekydream@qq.com
Copyright (c) 博尔塔拉融媒体中心媒体传播与认知重点实验室

The host sends the structured content as UTF-8 JSON on stdin and passes the output path as
the first argument, so the document layout lives here rather than in the plugin's JavaScript.
"""

import json
import sys

from docx import Document


def build(data):
    """Build the document from the host's payload."""
    document = Document()
    document.add_heading(str(data.get("title", "会议纪要")), level=0)
    for entry in data.get("meta", []):
        label, value = entry[0], entry[1]
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value))
    for section in data.get("sections", []):
        document.add_heading(str(section.get("title", "")), level=1)
        for item in section.get("items", []):
            document.add_paragraph(str(item), style="List Bullet")
    table = data.get("table")
    if table:
        document.add_heading(str(table.get("title", "")), level=1)
        headers = [str(header) for header in table.get("headers", [])]
        grid = document.add_table(rows=1, cols=max(1, len(headers)))
        grid.style = "Table Grid"
        for index, header in enumerate(headers):
            grid.rows[0].cells[index].text = header
        for row in table.get("rows", []):
            cells = grid.add_row().cells
            for index, value in enumerate(row[:len(headers)]):
                cells[index].text = str(value)
    transcript = data.get("transcript")
    if transcript:
        document.add_heading(str(transcript.get("title", "")), level=1)
        for line in transcript.get("lines", []):
            document.add_paragraph(str(line))
    footer = data.get("footer")
    if footer:
        document.add_paragraph(str(footer))
    return document


def main():
    """Read the payload, write the document beside the meeting's other exports."""
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "missing output path"}))
        return 2
    raw = sys.stdin.buffer.read().decode("utf-8")
    data = json.loads(raw)
    document = build(data)
    document.save(sys.argv[1])
    print(json.dumps({"ok": True, "path": sys.argv[1]}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
